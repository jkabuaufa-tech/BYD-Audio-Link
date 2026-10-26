#!/usr/bin/env python3
"""Ubah sheet cetak di master Excel menjadi template cetak HTML (cetak-template.js).

Template menyimpan tata letak persis seperti master: lebar kolom, tinggi baris,
sel gabungan (merge), garis, huruf, warna, perataan, format angka, dan area cetak.
Sel berisi rumus disimpan sebagai sel dinamis (alamat sel) yang diisi oleh
aplikasi (lihat CETAK_ISI di modul.js). Baris rincian item disimpan sebagai
"prototipe" per jenis baris (item, judul, subjudul, kelompok, kosong).

Pemakaian:
    python alat/konversi-cetak.py <folder-master>
Folder master berisi: Master_SPM.xlsm, BAP_format.xlsx (Subkon), BAP_SPSA.xlsx (Alat).
"""
import json
import pathlib
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

ROOT = pathlib.Path(__file__).resolve().parent.parent

# id template → file, sheet, blok rincian (baris awal-akhir) + baris prototipe tiap jenis
TEMPLATES = [
    dict(id='bahan_spm', file='spm', sheet='SPM', blok=(13, 112), proto={'item': 13}, isi=2),
    dict(id='bahan_bapm', file='spm', sheet='BAPM', blok=(26, 125), proto={'item': 26}, isi=2),
    dict(id='bahan_bap', file='spm', sheet='BAP'),
    dict(id='subkon_bapp', file='subkon', sheet='BAPP', blok=(34, 59),
         proto={'kelompok': 34, 'subjudul': 37, 'judul': 36, 'item': 38, 'kosong': 46}),
    dict(id='subkon_bap', file='subkon', sheet='BAP  1'),
    dict(id='subkon_opname', file='subkon', sheet='rincian ', blok=(11, 36),
         proto={'kelompok': 11, 'subjudul': 14, 'judul': 13, 'item': 15, 'kosong': 23}),
    dict(id='subkon_lampiran', file='subkon', sheet='lamp1', blok=(13, 38),
         proto={'kelompok': 13, 'subjudul': 16, 'judul': 15, 'item': 17, 'kosong': 25}),
    dict(id='subkon_bapr', file='subkon', sheet='BAPR ', blok=(32, 37), proto={'item': 33, 'pertama': 32}),
    dict(id='alat_bapa', file='alat', sheet='BAPA', blok=(30, 33),
         proto={'judul': 30, 'item': 31, 'subjudul': 33}, isi=1),
    dict(id='alat_bap', file='alat', sheet='BAP'),
    dict(id='alat_opname', file='alat', sheet='rincian ', blok=(11, 14),
         proto={'judul': 11, 'item': 12, 'subjudul': 14}),
    dict(id='alat_lampiran', file='alat', sheet='lamp', blok=(13, 15),
         proto={'judul': 13, 'item': 14}, isi=1),
    dict(id='alat_bapr', file='alat', sheet='BAPR ', blok=(32, 33), proto={'item': 33, 'pertama': 32}),
]

FILES = {'spm': 'Master_SPM.xlsm', 'subkon': 'BAP_format.xlsx', 'alat': 'BAP_SPSA.xlsx'}

BORDER = {
    'thin': '1px solid #000', 'hair': '1px solid #000', 'dotted': '1px dotted #000',
    'dashed': '1px dashed #000', 'mediumDashed': '2px dashed #000', 'dashDot': '1px dashed #000',
    'medium': '2px solid #000', 'thick': '3px solid #000', 'double': '3px double #000',
}


# Lebar digit maksimum (px) font default workbook — satuan lebar kolom Excel
MDW = {('arial', 10): 7, ('arial', 11): 8, ('arial', 12): 9, ('calibri', 11): 7, ('aptos narrow', 11): 7, ('aptos', 11): 8}


def px_col(w, mdw):
    if not w:
        return 64
    return int(((256 * w + int(128 / mdw)) / 256) * mdw)


def px_row(h):
    return round(h * 96 / 72, 1)


def color(c):
    try:
        if c is not None and c.type == 'rgb' and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb != '00000000':
            return '#' + c.rgb[2:]
    except Exception:
        pass
    return None


def cell_css(cell, edges):
    f, a = cell.font, cell.alignment
    css = []
    if f is not None:
        if f.name:
            css.append(f"font-family:'{f.name}',Arial,sans-serif")
        if f.sz:
            css.append(f'font-size:{float(f.sz):g}pt')
        if f.b:
            css.append('font-weight:bold')
        if f.i:
            css.append('font-style:italic')
        if f.u:
            css.append('text-decoration:underline')
        col = color(f.color)
        if col and col.lower() != '#000000':
            css.append(f'color:{col}')
    if a is not None:
        h = a.horizontal
        if h in ('center', 'centerContinuous'):
            css.append('text-align:center')
        elif h == 'right':
            css.append('text-align:right')
        elif h in ('justify', 'distributed'):
            css.append('text-align:justify')
        elif h == 'left':
            css.append('text-align:left')
        v = a.vertical
        css.append('vertical-align:' + {'top': 'top', 'center': 'middle', 'justify': 'top'}.get(v, 'bottom'))
        if a.wrap_text or h in ('justify', 'distributed') or v in ('justify', 'distributed'):
            css.append('white-space:pre-wrap')   # Excel melipat teks rata kanan-kiri walau "wrap" tidak dicentang
        if a.indent:
            css.append(f'padding-left:{int(a.indent) * 9 + 2}px')
    if cell.fill is not None and cell.fill.fill_type == 'solid':
        col = color(cell.fill.fgColor)
        if col and col.lower() != '#ffffff':
            css.append(f'background:{col}')
    for side in ('top', 'right', 'bottom', 'left'):
        b = getattr(edges[side].border, side)
        if b is not None and b.style in BORDER:
            css.append(f'border-{side}:{BORDER[b.style]}')
    return ';'.join(css)


def convert(cfg, wb_f, wb_v, styles):
    ws, wv = wb_f[cfg['sheet']], wb_v[cfg['sheet']]
    pa = ws.print_area if isinstance(ws.print_area, str) else ws.print_area[0]
    c1, r1, c2, r2 = range_boundaries(pa.split('!')[-1].replace('$', ''))
    # lebar/sembunyi kolom; satu entri bisa mencakup rentang kolom (min..max)
    lebar_kol, sembunyi_kol = {}, set()
    for dim in list(ws.column_dimensions.values()):
        lo, hi = dim.min or 0, dim.max or 0
        for c in range(lo, min(hi, c2) + 1):
            if dim.width:
                lebar_kol[c] = dim.width
            if dim.hidden:
                sembunyi_kol.add(c)
    cols_idx = [c for c in range(c1, c2 + 1) if c not in sembunyi_kol]
    default_w = ws.sheet_format.defaultColWidth or 8.43
    f0 = wb_f._fonts[0]
    mdw = MDW.get(((f0.name or '').lower(), int(f0.sz or 11)), 7)
    cols = []
    for c in cols_idx:
        cols.append(px_col(lebar_kol.get(c, default_w), mdw))
    default_h = ws.sheet_format.defaultRowHeight or 15

    merges = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        if m.max_row < r1 or m.min_row > r2 or m.max_col < c1 or m.min_col > c2:
            continue
        merges[(m.min_row, m.min_col)] = m
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    covered.add((r, c))

    def style_id(css):
        if css not in styles:
            styles[css] = len(styles)
        return styles[css]

    def row_json(r):
        rd = ws.row_dimensions[r]
        cells = []
        for i, c in enumerate(cols_idx):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            m = merges.get((r, c))
            cs = rs = 1
            edges = {'top': cell, 'left': cell, 'right': cell, 'bottom': cell}
            if m:
                vis = [x for x in cols_idx if m.min_col <= x <= min(m.max_col, c2)]
                cs = len(vis)
                rs = len([x for x in range(m.min_row, min(m.max_row, r2) + 1) if not ws.row_dimensions[x].hidden])
                edges['right'] = ws.cell(r, min(m.max_col, c2))
                edges['bottom'] = ws.cell(min(m.max_row, r2), c)
            item = {'i': i, 's': style_id(cell_css(cell, edges))}
            if cs > 1:
                item['cs'] = cs
            if rs > 1:
                item['rs'] = rs
            raw = cell.value
            val = wv.cell(r, c).value
            konstan = cell.data_type == 'f' and isinstance(raw, str) and re.fullmatch(r'=\s*"[^"]*"\s*', raw)
            if cell.data_type == 'f' and not konstan:
                item['a'] = cell.coordinate          # sel dinamis, diisi aplikasi
                if CONTOH and val not in (None, ''):
                    item['x'] = val if isinstance(val, (int, float)) else str(val)
            elif val not in (None, ''):
                item['v'] = val if isinstance(val, (int, float)) else str(val)
            fmt = cell.number_format
            if fmt and fmt != 'General':
                item['n'] = fmt
            cells.append(item)
        return {'h': px_row(rd.height if rd.height else default_h), 'r': r, 'c': cells}

    blok = cfg.get('blok')
    head, tail, protos = [], [], {}
    for r in range(r1, r2 + 1):
        if ws.row_dimensions[r].hidden and not (blok and blok[0] <= r <= blok[1]):
            continue
        if blok and blok[0] <= r <= blok[1]:
            continue
        (head if not blok or r < blok[0] else tail).append(row_json(r))
    for name, r in (cfg.get('proto') or {}).items():
        row = row_json(r)
        for c in row['c']:                      # isi baris rincian selalu dari data, bukan contoh master
            c.pop('v', None)
            c.pop('x', None)
        protos[name] = row

    logo = None
    imgs = sorted(((i.anchor._from.row + 1, i.anchor._from.col + 1) for i in ws._images if hasattr(i.anchor, '_from')))
    if imgs and imgs[0][0] <= r1 + 8:
        logo = {'r': imgs[0][0], 'i': max(0, min(cols_idx.index(imgs[0][1]) if imgs[0][1] in cols_idx else 0, len(cols) - 1))}

    m = ws.page_margins
    ps = ws.page_setup
    fit_page = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
    fit = {'page': fit_page, 'w': int(ps.fitToWidth if ps.fitToWidth is not None else 1),
           'h': int(ps.fitToHeight if ps.fitToHeight is not None else 1), 'scale': int(ps.scale or 100)}
    return {
        'fit': fit,
        'id': cfg['id'], 'sheet': cfg['sheet'].strip(),
        'orient': ws.page_setup.orientation or 'portrait',
        'margin': [m.top, m.right, m.bottom, m.left],
        'cols': cols, 'L': [get_column_letter(c) for c in cols_idx], 'head': head, 'tail': tail, 'proto': protos,
        'isi': cfg.get('isi', 0), 'logo': logo,
    }


CONTOH = False


def main():
    global CONTOH
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    CONTOH = '--contoh' in sys.argv   # sertakan nilai contoh dari master (hanya untuk uji, jangan dirilis)
    src = pathlib.Path(args[0] if args else '.')
    styles = {}
    out = {}
    cache = {}
    for cfg in TEMPLATES:
        f = cfg['file']
        if f not in cache:
            path = src / FILES[f]
            print('baca', path)
            cache[f] = (openpyxl.load_workbook(path, keep_links=False),
                        openpyxl.load_workbook(path, data_only=True, keep_links=False))
        out[cfg['id']] = convert(cfg, *cache[f], styles)
        t = out[cfg['id']]
        print(f"  {cfg['id']:16s} {len(t['cols'])} kolom, {len(t['head'])}+{len(t['tail'])} baris, proto {list(t['proto'])}")
    css = [None] * len(styles)
    for k, v in styles.items():
        css[v] = k
    js = ('/* Dibuat otomatis oleh alat/konversi-cetak.py dari master Excel — jangan diedit manual */\n'
          'window.CETAK_TPL = ' + json.dumps({'css': css, 'tpl': out}, ensure_ascii=False, separators=(',', ':')) + ';\n')
    dest = ROOT / 'cetak-template.js' if not CONTOH else pathlib.Path('cetak-template-contoh.js')
    dest.write_text(js, encoding='utf-8')
    print(f'{dest.name}: {len(js) / 1024:.0f} KB, {len(css)} gaya')


if __name__ == '__main__':
    main()
